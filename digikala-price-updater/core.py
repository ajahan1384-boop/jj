#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
منطق مشترک پردازش اکسل — بدون هیچ وابستگی به رابط کاربری (نه tkinter، نه
Flask) تا هم نسخه دسکتاپ (price_updater.py) و هم نسخه وب (web_app.py) از
همین یک تابع استفاده کنند.

هر عدد ستون J («قیمت فروش») شیت «داده ها» در ضریب داده‌شده ضرب می‌شود و
نتیجه تا ۴ رقم اعشار به سمت بالا (Ceiling) گرد می‌شود. بقیه‌ی فایل (شیت
راهنما، استایل‌ها، ستون‌های دیگر) دست‌نخورده باقی می‌ماند چون فایل ورودی
به‌جای ساخت از نو، مستقیماً ویرایش و ذخیره می‌شود.
"""

from __future__ import annotations

import os
import traceback
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_CEILING

try:
    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException
except ImportError:
    openpyxl = None
    InvalidFileException = Exception

SHEET_NAME = "داده ها"
PRICE_COLUMN_INDEX = 10  # ستون J
PRICE_COLUMN_LETTER = "J"
MAX_RECOMMENDED_FILES = 50
DECIMAL_PLACES = Decimal("0.0001")


@dataclass
class FileResult:
    input_path: str
    ok: bool
    message: str
    output_path: str = ""
    rows_updated: int = 0
    rows_skipped: int = 0


def _round_up_4(value: Decimal) -> Decimal:
    return value.quantize(DECIMAL_PLACES, rounding=ROUND_CEILING)


def process_one_file(input_path: str, output_dir: str, multiplier_str: str) -> FileResult:
    """یک فایل اکسل را می‌خواند، ستون J شیت «داده ها» را ضربدر ضریب می‌کند
    و تا ۴ رقم اعشار گرد به بالا می‌کند. همه خطاها اینجا گرفته می‌شوند تا
    کل برنامه هیچ‌وقت با خطای یک فایل متوقف/کرش نشود."""
    filename = os.path.basename(input_path)
    try:
        try:
            multiplier = Decimal(multiplier_str)
        except InvalidOperation:
            return FileResult(input_path, False, "ضریب واردشده یک عدد معتبر نیست.")

        if not os.path.isfile(input_path):
            return FileResult(input_path, False, "فایل پیدا نشد.")

        if openpyxl is None:
            return FileResult(input_path, False, "کتابخانه openpyxl نصب نیست.")

        try:
            wb = openpyxl.load_workbook(input_path, data_only=False)
        except InvalidFileException:
            return FileResult(input_path, False, "فرمت فایل معتبر نیست (باید xlsx باشد).")
        except PermissionError:
            return FileResult(input_path, False, "فایل توسط برنامه دیگری باز است (اکسل را ببندید).")
        except Exception as exc:  # noqa: BLE001
            return FileResult(input_path, False, f"خطا در باز کردن فایل: {exc}")

        if SHEET_NAME in wb.sheetnames:
            ws = wb[SHEET_NAME]
        else:
            # اگر نام شیت کمی متفاوت بود، شیتی که شبیه‌ترین نام را دارد پیدا می‌شود
            candidates = [s for s in wb.sheetnames if "داده" in s]
            if not candidates:
                return FileResult(
                    input_path, False,
                    f'شیت «{SHEET_NAME}» در فایل پیدا نشد. شیت‌های موجود: {", ".join(wb.sheetnames)}'
                )
            ws = wb[candidates[0]]

        rows_updated = 0
        rows_skipped = 0
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=PRICE_COLUMN_INDEX)
            value = cell.value
            if value is None:
                continue
            if isinstance(value, bool):
                rows_skipped += 1
                continue
            if isinstance(value, (int, float)):
                try:
                    new_value = _round_up_4(Decimal(str(value)) * multiplier)
                    cell.value = float(new_value)
                    rows_updated += 1
                except (InvalidOperation, ValueError):
                    rows_skipped += 1
            else:
                # مقدار متنی/فرمول در ستون قیمت -> برای جلوگیری از خرابی فایل دست نمی‌خورد
                rows_skipped += 1

        os.makedirs(output_dir, exist_ok=True)
        base, ext = os.path.splitext(filename)
        mult_label = multiplier_str.replace(".", "_").replace("/", "_")
        out_name = f"{base}_x{mult_label}.xlsx"
        out_path = os.path.join(output_dir, out_name)
        counter = 2
        while os.path.exists(out_path):
            out_path = os.path.join(output_dir, f"{base}_x{mult_label}_{counter}.xlsx")
            counter += 1

        try:
            wb.save(out_path)
        except PermissionError:
            return FileResult(input_path, False, "امکان ذخیره فایل خروجی نبود (دسترسی/فضای دیسک را بررسی کنید).")
        except Exception as exc:  # noqa: BLE001
            return FileResult(input_path, False, f"خطا در ذخیره فایل خروجی: {exc}")

        msg = f"{rows_updated} ردیف به‌روزرسانی شد"
        if rows_skipped:
            msg += f" ({rows_skipped} ردیف بدون تغییر رد شد)"
        return FileResult(input_path, True, msg, out_path, rows_updated, rows_skipped)

    except Exception as exc:  # noqa: BLE001  -- شبکه ایمنی نهایی؛ هرگز نباید کل برنامه کرش کند
        return FileResult(input_path, False, f"خطای غیرمنتظره: {exc}\n{traceback.format_exc(limit=2)}")
